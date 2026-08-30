from __future__ import annotations

import docx

from application.ports.document_parser_port import DocumentParserPort
from domain.entities import Document
from domain.enums import DocumentType


class DocxParser(DocumentParserPort):
    def supports(self, file_path: str) -> bool:
        return file_path.lower().endswith(".docx")

    def parse(self, file_path: str) -> Document:
        doc = docx.Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        table_lines: list[str] = []
        for table in doc.tables:
            for row in table.rows:
                table_lines.append(" | ".join(cell.text.strip() for cell in row.cells))

        raw_text = "\n".join(paragraphs)
        if table_lines:
            raw_text += "\n\n[Tableaux]\n" + "\n".join(table_lines)

        return Document(
            source_path=file_path,
            doc_type=DocumentType.DOCX,
            raw_text=raw_text,
            metadata={"paragraph_count": len(paragraphs), "table_count": len(doc.tables)},
        )
