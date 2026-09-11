from __future__ import annotations
from docx import Document as DocxDocument

from application.ports.document_parser_port import DocumentParserPort
from domain.entities import Document
from domain.enums import DocumentType


class DocxParser(DocumentParserPort):
    def parse(self, file_path: str) -> Document:
        doc = DocxDocument(file_path)
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return Document(
            source_path=file_path,
            doc_type=DocumentType.DOCX,
            raw_text="\n".join(parts),
            metadata={"paragraph_count": len(doc.paragraphs)},
        )
